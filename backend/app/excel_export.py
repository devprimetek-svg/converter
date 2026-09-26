"""Excel Export Module

Generates professionally formatted .xlsx workbooks from extracted parts catalogue data
using openpyxl according to Yamaha catalogue specifications.
"""

from __future__ import annotations

import io
import re
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.parts_extractor import clean_part_number, disambiguate_repeated_ref_numbers


def build_catalogue_code(model_code: str, fig_name: str, fig_no: str = "") -> str:
    """Build standardized catalogue code for a figure/assembly parent cell.
    Format: YAM_{MODEL_CODE}_{PARTS NAME} with space between words of parts name.
    e.g. YAM_BGPK_CYLINDER HEAD
    """
    mc = re.sub(r'[^A-Za-z0-9]+', '_', (model_code or "").strip()).strip('_').upper()
    if not mc:
        mc = "MODEL"
    # User rule: between words of parts name inside catalogue code, use space instead of underscore
    fn = re.sub(r'[^A-Za-z0-9]+', ' ', (fig_name or "").strip()).strip().upper()
    if not fn:
        if fig_no:
            padded = str(fig_no).zfill(2) if str(fig_no).isdigit() else str(fig_no)
            fn = f"FIG {padded}"
        else:
            fn = "PARTS"
    return f"YAM_{mc}_{fn}"


def build_image_record(model_code: str, fig_name: str, fig_no: str = "") -> str:
    """Build standardized parent image filename record.
    Matches naming standard: YAM_{MODEL_CODE}_{PARTS NAME}.jpeg
    e.g. YAM_BGPK_CYLINDER HEAD.jpeg
    """
    cat_code = build_catalogue_code(model_code, fig_name, fig_no)
    return f"{cat_code}.jpeg" if cat_code else ""


def build_model_name_record(model_code: str, model_name: str, fig_name: str) -> str:
    """Build the formatted model name record for a parent parts row.
    Format: YAMAHA {MODEL_CODE} {MODEL_NAME} Series {PARTS_NAME}
    e.g.    YAMAHA BGPJ LCX125 Series CYLINDER HEAD
    """
    mc = (model_code or "MODEL").strip().upper()
    mn = (model_name or "").strip()
    fn = re.sub(r"\s+", " ", (fig_name or "").strip()).upper()
    if not fn:
        fn = "PARTS"
    if mn:
        return f"YAMAHA {mc} {mn} Series {fn}"
    return f"YAMAHA {mc} Series {fn}"


def is_valid_quantity(val: Any) -> bool:
    """Return True if quantity is non-empty, non-zero, and not a blank placeholder."""
    if val is None:
        return False
    s = str(val).strip()
    if not s:
        return False
    if s.lower() in ("-", "none", "null", "0", "*"):
        return False
    return True


def generate_excel_workbook(
    rows: list[dict[str, Any]],
    model_columns: list[str],
    clean_parts: bool = False,
    sheet_title: str = "Parts List",
    model_code: Optional[str] = None,
) -> io.BytesIO:
    """Create an in-memory Excel workbook (.xlsx) from parts rows.

    Filtering rule:
    - If a specific model is being exported (len(model_columns) == 1), rows where that
      model's quantity is blank are excluded along with their remarks.
    - If multiple models exist, rows where all model quantities are blank are excluded.

    Parent cell rule:
    - For the figure number, parts name, and catalogue code, only the parent cell of the figure
      contains the value. Child rows under the same figure remain blank.
    - An extra 'Catalogue Code' column is added immediately after 'Parts Name'.
    """
    # Enforce blank quantity exclusion
    if len(model_columns) == 1:
        target_model = model_columns[0]
        rows = [r for r in rows if is_valid_quantity(r.get(target_model))]
        if sheet_title == "Parts List":
            sheet_title = f"Parts_{target_model}"[:31]
    elif len(model_columns) >= 2:
        rows = [r for r in rows if any(is_valid_quantity(r.get(m)) for m in model_columns)]

    # Disambiguate repeated ref_no within each figure (e.g. 1A, 1B, 2A, 2B etc.)
    rows = disambiguate_repeated_ref_numbers(rows)

    # Determine active model code for catalogue code generation
    active_model_code = (model_code or "").strip().upper()
    if not active_model_code:
        if len(model_columns) == 1:
            active_model_code = model_columns[0].strip().upper()
        elif len(model_columns) > 1:
            # Check if any row specifies a primary model_code
            for r in rows:
                if r.get("model_code"):
                    active_model_code = str(r["model_code"]).strip().upper()
                    break
            if not active_model_code:
                active_model_code = "_".join(m.strip().upper() for m in model_columns if m.strip())
        else:
            active_model_code = "MODEL"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    # Header styling
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")  # Yamaha Navy Blue
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data styling
    data_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border_side = Side(style="thin", color="D3D3D3")
    thin_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    # Base columns: Catalogue Code placed after Parts Name, Pic (parent image record) placed before Ref No.
    headers = [
        ("Page", "page", center_align),
        ("Fig No.", "fig_no", center_align),
        ("Parts Name", "fig_name", left_align),
        ("Catalogue Code", "catalogue_code", left_align),
        ("Model Name", "model_name", left_align),
        ("Pic", "pic", left_align),
        ("Ref No.", "ref_no", center_align),
        ("Part No.", "part_no", left_align),
        ("Description", "description", left_align),
    ]

    # Dynamic model columns
    for model in model_columns:
        headers.append((model, model, center_align))

    # Remarks column
    headers.append(("Remarks", "remarks", left_align))

    # Write header row (Row 1)
    ws.row_dimensions[1].height = 28.0
    for col_idx, (header_label, _, _) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data rows
    last_fig_key = None

    for row_idx, row_data in enumerate(rows, start=2):
        ws.row_dimensions[row_idx].height = 20.0

        # Identify figure grouping key
        raw_fig_no = str(row_data.get("parent_fig_no") or row_data.get("fig_no") or "").strip()
        raw_fig_name = str(row_data.get("parent_fig_name") or row_data.get("fig_name") or "").strip()
        if raw_fig_no or raw_fig_name:
            fig_key = (raw_fig_no, raw_fig_name)
        else:
            fig_key = ("page", str(row_data.get("page", "")))

        is_parent_cell = (fig_key != last_fig_key)
        if is_parent_cell:
            last_fig_key = fig_key
            curr_fig_no = raw_fig_no
            curr_fig_name = raw_fig_name
            curr_cat_code = row_data.get("catalogue_code") or build_catalogue_code(active_model_code, raw_fig_name, raw_fig_no)
            curr_pic = row_data.get("pic") or row_data.get("image") or (f"{curr_cat_code}.jpeg" if curr_cat_code else "")
            # Model name: prefer the pre-computed value from extraction; build fallback otherwise
            curr_model_name = row_data.get("model_name") or build_model_name_record(active_model_code, "", raw_fig_name)
        else:
            # Child cell: fig_no, fig_name, catalogue_code, model_name, and pic do NOT repeat — remain blank
            curr_fig_no = ""
            curr_fig_name = ""
            curr_cat_code = ""
            curr_model_name = ""
            curr_pic = ""

        for col_idx, (_, field_key, cell_align) in enumerate(headers, start=1):
            if field_key == "fig_no":
                val = curr_fig_no
            elif field_key == "fig_name":
                val = curr_fig_name
            elif field_key == "catalogue_code":
                val = curr_cat_code
            elif field_key == "model_name":
                val = curr_model_name
            elif field_key in ("pic", "image"):
                val = curr_pic
            elif field_key == "part_no":
                val = row_data.get("part_no", "")
                if clean_parts:
                    val = clean_part_number(str(val))
            else:
                val = row_data.get(field_key, "")

            cell = ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")
            cell.font = data_font
            cell.alignment = cell_align
            cell.border = thin_border
            # Enforce string data type so leading zeros aren't stripped
            cell.number_format = "@"

    # Freeze header row
    ws.freeze_panes = "A2"

    # Enable Autofilter
    if rows:
        max_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{max_col_letter}{len(rows) + 1}"

    # Calculate sensible auto-fit column widths
    for col_idx, (header_label, _, _) in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header_label)

        for row_idx in range(2, min(len(rows) + 2, 200)):  # sample up to 200 rows for speed
            cell_val = str(ws.cell(row=row_idx, column=col_idx).value or "")
            if len(cell_val) > max_len:
                max_len = len(cell_val)

        # Add buffer padding and clamp between min and max
        col_width = max(max_len + 4, 10)
        col_width = min(col_width, 48)  # max width cap
        ws.column_dimensions[col_letter].width = col_width

    # Save to memory buffer
    output_stream = io.BytesIO()
    wb.save(output_stream)
    output_stream.seek(0)
    return output_stream
