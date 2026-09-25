"""Tests for Product & Parts SEO Metadata Generator Module"""

import io
import openpyxl
import pytest

from app.parts_extractor import clean_part_number
from app.meta_generator import (
    build_product_title,
    build_meta_title,
    build_meta_short_description,
    build_meta_description,
    build_product_description,
    build_long_description,
    clean_no_commas,
    generate_main_part_metadata,
    generate_catalog_metadata,
    export_metadata_excel,
    export_metadata_csv,
    resolve_image_filename,
)


def test_resolve_image_filename():
    img_name = resolve_image_filename("CYLINDER HEAD", model_code="BGPK")
    assert img_name == "YAM_BGPK_CYLINDER HEAD.jpeg"


def test_no_commas_between_words():
    """Verify clean_no_commas strips all commas and normalizes spacing."""
    assert clean_no_commas("YAMAHA, BGPK, CYLINDER HEAD") == "YAMAHA BGPK CYLINDER HEAD"
    assert clean_no_commas("Hello, world, test, 123") == "Hello world test 123"


def test_product_title_in_caps_separate_from_meta_title_and_no_commas():
    """Verify product title is ALL CAPS, no commas, and meta title is separate (Title Case, no commas).
    Series is visible after Model in both short description and meta title.
    """
    pt = build_product_title("YAMAHA", "BGPK", "cylinder head", model="ray zr", series="series")
    assert pt == "YAMAHA BGPK RAY ZR SERIES CYLINDER HEAD"
    assert "," not in pt
    assert pt.isupper()

    mt = build_meta_title("YAMAHA", "BGPK", "cylinder head", model="ray zr", series="series")
    assert mt == "Yamaha Ray Zr Series BGPK Cylinder Head | IndiaSpare"
    assert "," not in mt
    assert mt != pt  # Separate from product title

    # Without model, series 'series' is not shown
    pt_no_model = build_product_title("YAMAHA", "BGPK", "cylinder head", model="", series="series")
    assert pt_no_model == "YAMAHA BGPK CYLINDER HEAD"
    mt_no_model = build_meta_title("YAMAHA", "BGPK", "cylinder head", model="", series="series")
    assert mt_no_model == "Yamaha BGPK Cylinder Head | IndiaSpare"


def test_meta_short_description_no_commas():
    """Verify meta short description ends with IndiaSpare and has no commas."""
    sd = build_meta_short_description("YAMAHA", "BGPK", "CYLINDER HEAD", model="RAY ZR")
    assert sd == "YAMAHA BGPK RAY ZR CYLINDER HEAD IndiaSpare"
    assert "," not in sd
    assert sd.endswith("IndiaSpare")


def test_meta_description_character_length_151_to_158_without_caps_no_commas():
    """Verify meta descriptions are strictly bounded between 151 and 158 characters, without caps except IndiaSpare, and zero commas."""
    sample_parts = [
        "CYLINDER HEAD",
        "CRANKSHAFT & PISTON",
        "VALVE",
        "AIR SHROUD & FAN",
        "OIL PUMP",
        "INTAKE",
        "EXHAUST",
        "CRANKCASE",
        "CRANKCASE COVER 1",
        "STARTER CLUTCH",
        "CLUTCH",
        "TRANSMISSION",
        "FRAME & SUBFRAME MOUNTING HARDWARE",
    ]

    for part in sample_parts:
        for m in ["", "RAY ZR"]:
            desc = build_meta_description(
                brand="Yamaha",
                model_code="BGPK",
                part_name=part,
                model=m,
                series="series",
            )
            assert 151 <= len(desc) <= 158, f"Part '{part}' (m='{m}') desc length {len(desc)} not in [151, 158]: '{desc}'"
            assert "," not in desc, f"Comma found in meta description: '{desc}'"
            assert "IndiaSpare" in desc, f"'IndiaSpare' missing from meta description: '{desc}'"
            assert "indiaspare" not in desc, f"Lowercase 'indiaspare' found in meta description: '{desc}'"
            assert "INDIASPARE" not in desc, f"Uppercase 'INDIASPARE' found in meta description: '{desc}'"
            rest = desc.replace("IndiaSpare", "")
            assert rest.islower(), f"Meta description outside 'IndiaSpare' must be without caps: '{desc}'"


def test_product_description_words_120_to_140_no_commas():
    """Verify product descriptions are strictly bounded between 120 and 140 words, and zero commas."""
    sample_parts = [
        "CYLINDER HEAD",
        "CRANKSHAFT & PISTON",
        "VALVE",
        "AIR SHROUD & FAN",
        "OIL PUMP",
        "TRANSMISSION",
    ]

    for part in sample_parts:
        for m in ["", "RAY ZR"]:
            pdesc = build_product_description(
                brand="YAMAHA",
                model_code="BGPK",
                part_name=part,
                model=m,
                series="series",
            )
            words = pdesc.split()
            word_count = len(words)
            assert 120 <= word_count <= 140, f"Part '{part}' word count {word_count} not in [120, 140]: '{pdesc}'"
            assert "," not in pdesc, f"Comma found in product description: '{pdesc}'"
            assert "IndiaSpare" in pdesc


def test_generate_main_part_metadata_blank_by_default():
    """Verify that by default (blank_descriptions=True), descriptions are blank."""
    figure = {"fig_no": "1", "fig_name": "CYLINDER HEAD", "first_page": 7}
    meta = generate_main_part_metadata(
        figure=figure,
        model_code="BGPK",
        brand="YAMAHA",
        model="",
        series="series",
    )

    # 1. Product Title: in ALL CAPS, no commas
    assert meta["product_title"] == "YAMAHA BGPK CYLINDER HEAD"
    assert "," not in meta["product_title"]
    assert meta["product_title"].isupper()

    # 2. Meta Title: separate from product title, Title Case, no commas
    assert meta["meta_title"] == "Yamaha BGPK Cylinder Head | IndiaSpare"
    assert "," not in meta["meta_title"]

    # 3. Meta Short description: removed from entire app
    assert "meta_short_description" not in meta

    # 4. Meta description & Product description: blank by default during scan
    assert meta["meta_description"] == ""
    assert meta["meta_long_description"] == ""
    assert meta["meta_desc_chars"] == 0
    assert meta["product_description"] == ""
    assert meta["product_desc_words"] == 0

    # 5. Image filename & fields
    assert meta["image_filename"] == "YAM_BGPK_CYLINDER HEAD.jpeg"
    assert meta["brand"] == "YAMAHA"
    assert meta["model"] == ""
    assert meta["series"] == "series"


def test_generate_main_part_with_descriptions_explicit():
    """Verify that when blank_descriptions=False, descriptions are populated and bounded."""
    figure = {"fig_no": "1", "fig_name": "CYLINDER HEAD", "first_page": 7}
    meta = generate_main_part_metadata(
        figure=figure,
        model_code="BGPK",
        brand="YAMAHA",
        model="",
        series="series",
        blank_descriptions=False,
    )

    assert 151 <= len(meta["meta_description"]) <= 158
    assert "," not in meta["meta_description"]
    assert "IndiaSpare" in meta["meta_description"]
    assert meta["meta_desc_chars"] == len(meta["meta_description"])

    p_words = len(meta["product_description"].split())
    assert 120 <= p_words <= 140
    assert "," not in meta["product_description"]
    assert meta["product_desc_words"] == p_words


def test_generate_main_part_with_custom_model_and_series():
    """Verify custom model name and series are included in title and descriptions."""
    figure = {"fig_no": "2", "fig_name": "CRANKSHAFT & PISTON", "first_page": 8}
    meta = generate_main_part_metadata(
        figure=figure,
        model_code="BGPK",
        brand="YAMAHA",
        model="RAY ZR",
        series="STREET RALLY",
        blank_descriptions=False,
    )

    # Title with series after model in ALL CAPS, no commas
    assert meta["product_title"] == "YAMAHA BGPK RAY ZR STREET RALLY CRANKSHAFT & PISTON"
    assert "," not in meta["product_title"]

    # Separate Meta Title (series visible after model, model before model code)
    assert meta["meta_title"] == "Yamaha Ray Zr Street Rally BGPK Crankshaft & Piston | IndiaSpare"
    assert "," not in meta["meta_title"]

    # Short desc removed from entire app
    assert "meta_short_description" not in meta

    # Meta desc strictly 151-158 chars without caps except 'IndiaSpare', no commas
    assert 151 <= len(meta["meta_description"]) <= 158
    assert "," not in meta["meta_description"]
    assert "IndiaSpare" in meta["meta_description"]
    assert "indiaspare" not in meta["meta_description"]
    assert meta["meta_description"].replace("IndiaSpare", "").islower()

    # Product desc strictly 120-140 words, no commas
    w_count = len(meta["product_description"].split())
    assert 120 <= w_count <= 140
    assert "," not in meta["product_description"]
    assert "IndiaSpare" in meta["product_description"]


def test_main_parts_only_catalog_extraction():
    """Verify catalog metadata contains ONLY the main parts (figures), and descriptions are blank by default."""
    rows = [
        {"fig_no": "1", "fig_name": "CYLINDER HEAD", "ref_no": "1", "part_no": "BGP-E1111-00", "description": "HEAD, CYLINDER 1", "page": 7},
        {"fig_no": "1", "fig_name": "CYLINDER HEAD", "ref_no": "2", "part_no": "90430-06817", "description": "GASKET", "page": 7},
        {"fig_no": "1", "fig_name": "CYLINDER HEAD", "ref_no": "3", "part_no": "95022-06010", "description": "BOLT, FLANGE", "page": 7},
        {"fig_no": "2", "fig_name": "CRANKSHAFT & PISTON", "ref_no": "1", "part_no": "BGP-E1400-00", "description": "CRANKSHAFT ASSY", "page": 8},
        {"fig_no": "2", "fig_name": "CRANKSHAFT & PISTON", "ref_no": "2", "part_no": "BGP-E1631-00", "description": "PISTON", "page": 8},
    ]

    # Default: blank descriptions
    meta_items = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="YAMAHA",
        model="",
        series="series",
        model_code="BGPK",
        main_parts_only=True,
    )

    assert len(meta_items) == 2
    assert meta_items[0]["part_name"] == "CYLINDER HEAD"
    assert meta_items[0]["product_title"] == "YAMAHA BGPK CYLINDER HEAD"
    assert meta_items[0]["meta_title"] == "Yamaha BGPK Cylinder Head | IndiaSpare"
    assert meta_items[0]["meta_description"] == ""
    assert meta_items[0]["product_description"] == ""

    # With blank_descriptions=False
    meta_filled = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="YAMAHA",
        model="",
        series="series",
        model_code="BGPK",
        main_parts_only=True,
        blank_descriptions=False,
    )
    assert 151 <= len(meta_filled[0]["meta_description"]) <= 158
    assert 120 <= len(meta_filled[0]["product_description"].split()) <= 140


def test_export_metadata_excel_and_csv():
    figure = {"fig_no": "1", "fig_name": "CYLINDER HEAD", "first_page": 7}
    meta_item = generate_main_part_metadata(
        figure=figure,
        model_code="BGPK",
        brand="YAMAHA",
        model="",
        series="series",
        blank_descriptions=False,
    )
    meta_items = [meta_item]

    # Test Excel generation
    xlsx_buf = export_metadata_excel(meta_items, brand="YAMAHA")
    assert xlsx_buf.getbuffer().nbytes > 0
    wb = openpyxl.load_workbook(xlsx_buf)
    ws = wb.active
    assert ws.title == "Product Metadata"
    assert ws.cell(row=1, column=1).value == "Page"
    assert ws.cell(row=1, column=2).value == "Fig No."
    assert ws.cell(row=1, column=3).value == "Catalog Name"
    assert ws.cell(row=1, column=4).value == "Catalogue Code"
    assert ws.cell(row=1, column=5).value == "Ref No."
    assert ws.cell(row=1, column=6).value == "Part No."
    assert ws.cell(row=1, column=7).value == "Clean Part No."
    assert ws.cell(row=1, column=8).value == "Description"
    assert ws.cell(row=1, column=9).value == "Remarks"
    assert ws.cell(row=1, column=10).value == "Brand"
    assert ws.cell(row=1, column=14).value == "Pic"
    assert ws.cell(row=1, column=15).value == "Short Description"
    assert ws.cell(row=1, column=16).value == "Meta Title"
    assert ws.cell(row=2, column=1).value == "7"
    assert ws.cell(row=2, column=2).value == "1"
    assert ws.cell(row=2, column=3).value == "CYLINDER HEAD"
    assert ws.cell(row=2, column=4).value == "YAM_BGPK_CYLINDER_HEAD"
    assert ws.cell(row=2, column=14).value == "YAM_BGPK_CYLINDER HEAD.jpeg"
    assert ws.cell(row=2, column=15).value == "YAMAHA BGPK CYLINDER HEAD"
    assert ws.cell(row=2, column=16).value == "Yamaha BGPK Cylinder Head | IndiaSpare"
    assert 151 <= len(ws.cell(row=2, column=17).value) <= 158
    assert 120 <= len(ws.cell(row=2, column=19).value.split()) <= 140

    # Test CSV generation
    csv_buf = export_metadata_csv(meta_items)
    assert csv_buf.getbuffer().nbytes > 0
    csv_content = csv_buf.getvalue().decode("utf-8-sig")
    assert "Short Description" in csv_content
    assert "Catalog Name" in csv_content
    assert "Pic" in csv_content
    assert "Meta Title" in csv_content
    assert "Meta Short Description" not in csv_content
    assert "Meta Description (151-158 Chars)" in csv_content
    assert "Product Description (120-140 Words)" in csv_content
    assert "YAMAHA BGPK CYLINDER HEAD" in csv_content
    assert "Yamaha BGPK Cylinder Head | IndiaSpare" in csv_content
    assert "INDIASPARE" not in csv_content


def test_export_both_extracted_and_seo_columns():
    """Verify Excel and CSV exports contain BOTH extracted catalogue columns and SEO metadata columns."""
    rows = [
        {
            "fig_no": "1",
            "fig_name": "CYLINDER HEAD",
            "ref_no": "1",
            "part_no": "B7J-E1102-00",
            "description": "CYLINDER HEAD ASSY",
            "BGPK": "1",
            "remarks": "UR FOR VRC1",
            "page": 7,
        }
    ]
    meta_items = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="YAMAHA",
        model="FZ",
        series="series",
        model_code="BGPK",
        main_parts_only=False,
        blank_descriptions=False,
    )

    # 1. Excel export with both extracted columns and SEO columns
    xlsx_buf = export_metadata_excel(
        meta_items,
        brand="YAMAHA",
        model_columns=["BGPK"],
        raw_rows=rows,
    )
    wb = openpyxl.load_workbook(xlsx_buf)
    ws = wb["Product Metadata"]

    header_vals = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert "Fig No." in header_vals
    assert "Catalog Name" in header_vals
    assert "Brand" in header_vals
    assert "Pic" in header_vals
    assert "Short Description" in header_vals
    assert "Meta Title" in header_vals
    assert "Ref No." in header_vals
    assert "Part No." in header_vals
    assert "Clean Part No." in header_vals
    assert "Qty (BGPK)" in header_vals
    assert "Remarks" in header_vals

    row2_vals = [str(ws.cell(row=2, column=c).value or "") for c in range(1, ws.max_column + 1)]
    assert "B7J-E1102-00" in row2_vals
    assert clean_part_number("B7J-E1102-00") in row2_vals
    assert "UR FOR VRC1" in row2_vals

    # Verify STRICTLY SINGLE SHEET ONLY: No separate sheet for child cells per user instruction
    assert len(wb.sheetnames) == 1
    assert wb.sheetnames == ["Product Metadata"]
    assert "Extracted Parts Catalogue" not in wb.sheetnames
    # Verify Excel AutoFilter is enabled so user can filter child cells directly in the single sheet
    assert ws.auto_filter.ref is not None

    # 2. CSV export with both extracted columns and SEO columns
    csv_buf = export_metadata_csv(meta_items, model_columns=["BGPK"])
    csv_text = csv_buf.getvalue().decode("utf-8-sig")
    assert "Ref No." in csv_text
    assert "Part No." in csv_text
    assert "Clean Part No." in csv_text
    assert "Qty (BGPK)" in csv_text
    assert "B7J-E1102-00" in csv_text
    assert "UR FOR VRC1" in csv_text
    assert "Short Description" in csv_text


def test_column_selection_filtering():
    """Verify selecting specific columns filters exported Excel and CSV to only those columns."""
    figure = {"fig_no": "1", "fig_name": "CYLINDER HEAD", "first_page": 7}
    meta_item = generate_main_part_metadata(
        figure=figure,
        model_code="BGPK",
        brand="YAMAHA",
        model="FZ-S",
        series="series",
        blank_descriptions=False,
    )
    meta_items = [meta_item]

    # Select only 3 columns
    selected = ["Fig No.", "Catalog Name", "Meta Title"]

    # Excel export
    xlsx_buf = export_metadata_excel(meta_items, selected_columns=selected)
    wb = openpyxl.load_workbook(xlsx_buf)
    ws = wb["Product Metadata"]
    exported_headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert exported_headers == ["Fig No.", "Catalog Name", "Meta Title"]
    assert ws.cell(row=2, column=1).value == "1"
    assert ws.cell(row=2, column=2).value == "CYLINDER HEAD"
    assert ws.cell(row=2, column=3).value == "Yamaha Fz-S Series BGPK Cylinder Head | IndiaSpare"

    # CSV export
    csv_buf = export_metadata_csv(meta_items, selected_columns=selected)
    csv_text = csv_buf.getvalue().decode("utf-8-sig")
    lines = [line.strip() for line in csv_text.strip().split("\r\n") if line.strip()]
    assert lines[0] == "Fig No.,Catalog Name,Meta Title"
    assert "Yamaha Fz-S Series BGPK Cylinder Head | IndiaSpare" in lines[1]
    assert "Short Description" not in csv_text


def test_model_output_always_in_all_caps():
    """Verify that even when user inputs lowercase or mixed case model name, output is in ALL CAPS."""
    figure = {"fig_no": "1", "fig_name": "CYLINDER HEAD", "first_page": 7}
    meta = generate_main_part_metadata(
        figure=figure,
        model_code="bgpk",
        brand="yamaha",
        model="ray zr 125 fi",
        series="street rally",
        blank_descriptions=True,
    )
    assert meta["model"] == "RAY ZR 125 FI"
    assert meta["product_title"] == "YAMAHA BGPK RAY ZR 125 FI STREET RALLY CYLINDER HEAD"

    # In catalog metadata
    rows = [{"fig_no": "1", "fig_name": "CYLINDER HEAD", "page": 1}]
    items = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="yamaha",
        model="fz-s v3",
        series="series",
        model_code="bgpk",
        main_parts_only=True,
    )
    assert items[0]["model"] == "FZ-S V3"


def test_catalogue_columns_arranged_on_left_before_metadata_columns():
    """Verify extracted catalogue columns appear strictly on the left followed by SEO & metadata columns."""
    rows = [
        {
            "fig_no": "1",
            "fig_name": "CYLINDER HEAD",
            "ref_no": "1",
            "part_no": "B7J-E1102-00",
            "description": "CYLINDER HEAD ASSY",
            "BGPK": "1",
            "BGPL": "2",
            "remarks": "UR FOR VRC1",
            "page": 7,
        }
    ]
    meta_items = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK", "BGPL"],
        brand="YAMAHA",
        model="FZ-S",
        series="series",
        model_code="BGPK_BGPL",
        main_parts_only=False,
        blank_descriptions=False,
    )

    # 1. Excel Export
    xlsx_buf = export_metadata_excel(
        meta_items,
        brand="YAMAHA",
        model_columns=["BGPK", "BGPL"],
        raw_rows=rows,
    )
    wb = openpyxl.load_workbook(xlsx_buf)
    ws = wb["Product Metadata"]
    excel_headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    # Expected order: Extracted parts catalogue columns on LEFT
    expected_left_catalogue = [
        "Page",
        "Fig No.",
        "Catalog Name",
        "Catalogue Code",
        "Ref No.",
        "Part No.",
        "Clean Part No.",
        "Description",
        "Qty (BGPK)",
        "Qty (BGPL)",
        "Remarks",
    ]
    # Expected order: SEO & Metadata columns on RIGHT
    expected_right_metadata = [
        "Brand",
        "Model Code",
        "Model",
        "Series",
        "Pic",
        "Short Description",
        "Meta Title",
        "Meta Description (151-158 Chars)",
        "Meta Desc Chars",
        "Product Description (120-140 Words)",
        "Product Desc Words",
    ]

    assert excel_headers == expected_left_catalogue + expected_right_metadata

    # 2. CSV Export
    csv_buf = export_metadata_csv(meta_items, model_columns=["BGPK", "BGPL"])
    csv_lines = csv_buf.getvalue().decode("utf-8-sig").splitlines()
    csv_headers = csv_lines[0].split(",")

    assert csv_headers == expected_left_catalogue + expected_right_metadata


def test_child_cells_descriptions_blank_by_default_and_populated_when_prompted():
    """Verify child cells have blank descriptions by default, but receive descriptions when prompt mentions child cells."""
    rows = [
        # Parent row
        {
            "fig_no": "1",
            "fig_name": "CYLINDER HEAD",
            "ref_no": "1",
            "part_no": "B7J-E1102-00",
            "description": "CYLINDER HEAD ASSY",
            "BGPK": "1",
            "page": 7,
        },
        # Child row 1
        {
            "fig_no": "1",
            "fig_name": "CYLINDER HEAD",
            "ref_no": "2",
            "part_no": "90105-06836",
            "description": "BOLT FLANGE",
            "BGPK": "4",
            "page": 7,
        },
        # Child row 2
        {
            "fig_no": "1",
            "fig_name": "CYLINDER HEAD",
            "ref_no": "3",
            "part_no": "94700-00879",
            "description": "PLUG SPARK",
            "BGPK": "1",
            "page": 7,
        },
    ]

    # CASE 1: Standard generation (no child cells mentioned in prompt)
    items_default = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="YAMAHA",
        model="FZ-S",
        series="series",
        model_code="BGPK",
        main_parts_only=False,
        blank_descriptions=False,
        user_prompt="Focus on OEM build quality and durability",
    )
    assert len(items_default) == 3

    # Parent item
    p = items_default[0]
    assert p["is_parent"] is True
    assert p["cell_type"] == "Parent"
    assert p["fig_no"] == "1"
    assert p["part_name"] == "CYLINDER HEAD"
    assert p["catalogue_code"] == "YAM_BGPK_CYLINDER_HEAD"
    assert len(p["meta_description"]) >= 151
    assert len(p["product_description"].split()) >= 120

    # Child item 1
    c1 = items_default[1]
    assert c1["is_parent"] is False
    assert c1["cell_type"] == "Child"
    assert c1["fig_no"] == ""  # Blank per user rule
    assert c1["part_name"] == ""  # Blank per user rule
    assert c1["catalogue_code"] == ""  # Blank per user rule
    assert c1["description"] == "BOLT FLANGE"
    assert c1["ref_no"] == "2"
    assert c1["part_no"] == "90105-06836"
    assert c1["meta_description"] == ""  # Strictly blank per user rule!
    assert c1["product_description"] == ""  # Strictly blank per user rule!

    # Child item 2
    c2 = items_default[2]
    assert c2["is_parent"] is False
    assert c2["fig_no"] == ""
    assert c2["part_name"] == ""
    assert c2["catalogue_code"] == ""
    assert c2["meta_description"] == ""
    assert c2["product_description"] == ""

    # CASE 2: User explicitly mentions child cells in prompt
    items_with_child_prompt = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="YAMAHA",
        model="FZ-S",
        series="series",
        model_code="BGPK",
        main_parts_only=False,
        blank_descriptions=False,
        user_prompt="Generate descriptions for all parts including child cells with OEM precision",
    )
    assert len(items_with_child_prompt) == 3

    # Parent still has full descriptions
    p_prompt = items_with_child_prompt[0]
    assert p_prompt["is_parent"] is True
    assert len(p_prompt["meta_description"]) >= 151

    # Child item 1 NOW has generated descriptions because child cells were referenced!
    c1_prompt = items_with_child_prompt[1]
    assert c1_prompt["is_parent"] is False
    assert c1_prompt["fig_no"] == ""
    assert c1_prompt["part_name"] == ""
    assert c1_prompt["catalogue_code"] == ""
    assert len(c1_prompt["meta_description"]) >= 151
    assert len(c1_prompt["product_description"].split()) >= 120


def test_export_single_sheet_with_autofilter_for_child_cells():
    """Verify export is STRICTLY a single sheet and AutoFilter is applied across all columns for child filtering."""
    rows = [
        {"fig_no": "1", "fig_name": "CYLINDER HEAD", "ref_no": "1", "part_no": "B7J-E1102-00", "description": "CYLINDER HEAD ASSY", "BGPK": "1", "page": 7},
        {"fig_no": "1", "fig_name": "CYLINDER HEAD", "ref_no": "2", "part_no": "90105-06836", "description": "BOLT FLANGE", "BGPK": "4", "page": 7},
    ]
    meta_items = generate_catalog_metadata(
        rows=rows,
        model_columns=["BGPK"],
        brand="YAMAHA",
        model="RAY ZR",
        series="series",
        model_code="BGPK",
        main_parts_only=False,
        blank_descriptions=False,
    )

    xlsx_buf = export_metadata_excel(
        meta_items,
        brand="YAMAHA",
        model_columns=["BGPK"],
        raw_rows=rows,
    )
    wb = openpyxl.load_workbook(xlsx_buf)

    # Strictly 1 sheet
    assert len(wb.sheetnames) == 1
    assert wb.sheetnames == ["Product Metadata"]

    ws = wb["Product Metadata"]
    # AutoFilter is present so user can filter parent/child cells
    assert ws.auto_filter.ref is not None
    assert "A1:" in ws.auto_filter.ref

    # Row 2 (Parent) has Fig No and Catalogue Code populated
    assert ws.cell(row=2, column=2).value == "1"
    assert ws.cell(row=2, column=3).value == "CYLINDER HEAD"
    assert ws.cell(row=2, column=4).value == "YAM_BGPK_CYLINDER_HEAD"

    # Row 3 (Child) has Fig No, Catalog Name, Catalogue Code BLANK
    assert (ws.cell(row=3, column=2).value or "") == ""
    assert (ws.cell(row=3, column=3).value or "") == ""
    assert (ws.cell(row=3, column=4).value or "") == ""
    assert str(ws.cell(row=3, column=5).value) == "2"  # Ref No
    assert str(ws.cell(row=3, column=6).value) == "90105-06836"  # Part No
    assert str(ws.cell(row=3, column=8).value) == "BOLT FLANGE"  # Description



