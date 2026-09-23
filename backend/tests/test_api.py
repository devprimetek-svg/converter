"""API endpoint integration tests."""

import io
import time
from fastapi.testclient import TestClient

from app.main import app
from tests.test_parts_extractor import _create_synthetic_catalogue_pdf

client = TestClient(app)


def test_health_endpoint():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_extract_and_export_flow():
    # 1. Create synthetic PDF
    pdf_buf = _create_synthetic_catalogue_pdf()
    files = {"file": ("test_catalogue.pdf", pdf_buf.getvalue(), "application/pdf")}

    # 2. Upload to /api/extract
    response = client.post("/api/extract", files=files)
    assert response.status_code == 200
    data = response.json()
    assert "job_id" in data
    job_id = data["job_id"]

    # 3. Poll /api/extract/status/{job_id} until completed
    max_wait = 10
    start_time = time.time()
    status_data = None

    while time.time() - start_time < max_wait:
        status_res = client.get(f"/api/extract/status/{job_id}")
        assert status_res.status_code == 200
        status_data = status_res.json()
        if status_data["status"] in ("completed", "error"):
            break
        time.sleep(0.2)

    assert status_data is not None
    assert status_data["status"] == "completed"
    assert status_data["total_rows"] == 3
    assert "BGPK" in status_data["model_columns"]
    assert len(status_data["rows"]) == 3

    # 4. Test /api/export with clean part numbers
    export_payload = {
        "filename": "test_catalogue",
        "clean_part_numbers": True,
        "model_columns": status_data["model_columns"],
        "rows": status_data["rows"],
    }
    export_res = client.post("/api/export", json=export_payload)
    assert export_res.status_code == 200
    assert export_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "test_catalogue_Parts.xlsx" in export_res.headers.get("content-disposition", "")
    assert export_res.content[:4] == b"PK\x03\x04"


def test_extract_invalid_extension():
    files = {"file": ("test.txt", b"not a pdf", "text/plain")}
    response = client.post("/api/extract", files=files)
    assert response.status_code == 400
    assert "PDF files (.pdf) are supported" in response.json()["detail"]


def test_multi_model_export():
    """Verify exporting specific model code downloads filtered file."""
    rows = [
        {"page": 1, "fig_no": "1", "fig_name": "HEAD", "ref_no": "1", "part_no": "95022-06010", "description": "BOLT", "BGPJ": "1", "BGPL": "1", "remarks": ""},
        {"page": 1, "fig_no": "1", "fig_name": "HEAD", "ref_no": "2", "part_no": "90430-06817", "description": "GASKET", "BGPJ": "2", "BGPL": "", "remarks": ""},
        {"page": 1, "fig_no": "1", "fig_name": "HEAD", "ref_no": "3", "part_no": "95612-08618", "description": "STUD", "BGPJ": "", "BGPL": "1", "remarks": ""},
    ]

    # Export for BGPJ only
    res_bgpj = client.post(
        "/api/export",
        json={
            "filename": "Catalogue_2026",
            "clean_part_numbers": False,
            "model_columns": ["BGPJ", "BGPL"],
            "target_model": "BGPJ",
            "rows": rows,
        },
    )
    assert res_bgpj.status_code == 200
    assert "Catalogue_2026_BGPJ_Parts.xlsx" in res_bgpj.headers.get("content-disposition", "")

    # Export for BGPL only
    res_bgpl = client.post(
        "/api/export",
        json={
            "filename": "Catalogue_2026",
            "clean_part_numbers": False,
            "model_columns": ["BGPJ", "BGPL"],
            "target_model": "BGPL",
            "rows": rows,
        },
    )
    assert res_bgpl.status_code == 200
    assert "Catalogue_2026_BGPL_Parts.xlsx" in res_bgpl.headers.get("content-disposition", "")


def test_blank_quantity_row_and_remarks_omitted():
    """Verify that when a model code quantity is blank, the entire row and remarks are excluded from Excel."""
    import openpyxl

    rows = [
        {
            "page": 1,
            "fig_no": "1",
            "fig_name": "CYLINDER",
            "ref_no": "1",
            "part_no": "95022-06010",
            "description": "BOLT, FLANGE",
            "BGPK": "2",
            "BGPL": "-",
            "remarks": "UR FOR BGPK ONLY",
        },
        {
            "page": 1,
            "fig_no": "1",
            "fig_name": "CYLINDER",
            "ref_no": "2",
            "part_no": "90430-06817",
            "description": "GASKET",
            "BGPK": "",
            "BGPL": "1",
            "remarks": "UR FOR BGPL ONLY",
        },
        {
            "page": 1,
            "fig_no": "1",
            "fig_name": "CYLINDER",
            "ref_no": "3",
            "part_no": "99999-00000",
            "description": "UNASSIGNED PART",
            "BGPK": "",
            "BGPL": "",
            "remarks": "NEITHER MODEL",
        },
    ]

    # 1. Export BGPK only
    res_bgpk = client.post(
        "/api/export",
        json={
            "filename": "Test_MultiModel",
            "clean_part_numbers": False,
            "model_columns": ["BGPK", "BGPL"],
            "target_model": "BGPK",
            "rows": rows,
        },
    )
    assert res_bgpk.status_code == 200
    wb_bgpk = openpyxl.load_workbook(io.BytesIO(res_bgpk.content))
    ws_bgpk = wb_bgpk.active
    # Collect all cell text across all rows
    all_bgpk_texts = [cell.value for r in ws_bgpk.iter_rows() for cell in r if cell.value is not None]

    # Row 1 must be present (BGPK has qty 2)
    assert "95022-06010" in all_bgpk_texts
    assert "UR FOR BGPK ONLY" in all_bgpk_texts

    # Row 2 (BGPK is blank) must NOT be present (part no and remarks excluded)
    assert "90430-06817" not in all_bgpk_texts
    assert "UR FOR BGPL ONLY" not in all_bgpk_texts

    # Row 3 (both blank) must NOT be present
    assert "99999-00000" not in all_bgpk_texts
    assert "NEITHER MODEL" not in all_bgpk_texts

    # 2. Export BGPL only
    res_bgpl = client.post(
        "/api/export",
        json={
            "filename": "Test_MultiModel",
            "clean_part_numbers": False,
            "model_columns": ["BGPK", "BGPL"],
            "target_model": "BGPL",
            "rows": rows,
        },
    )
    assert res_bgpl.status_code == 200
    wb_bgpl = openpyxl.load_workbook(io.BytesIO(res_bgpl.content))
    ws_bgpl = wb_bgpl.active
    all_bgpl_texts = [cell.value for r in ws_bgpl.iter_rows() for cell in r if cell.value is not None]

    # Row 2 must be present (BGPL has qty 1)
    assert "90430-06817" in all_bgpl_texts
    assert "UR FOR BGPL ONLY" in all_bgpl_texts

    # Row 1 (BGPL is "-") must NOT be present
    assert "95022-06010" not in all_bgpl_texts
    assert "UR FOR BGPK ONLY" not in all_bgpl_texts

    # 3. Export ALL models
    res_all = client.post(
        "/api/export",
        json={
            "filename": "Test_MultiModel",
            "clean_part_numbers": False,
            "model_columns": ["BGPK", "BGPL"],
            "rows": rows,
        },
    )
    assert res_all.status_code == 200
    wb_all = openpyxl.load_workbook(io.BytesIO(res_all.content))
    ws_all = wb_all.active
    all_texts = [cell.value for r in ws_all.iter_rows() for cell in r if cell.value is not None]

    # Row 1 and Row 2 have at least one valid qty, so present
    assert "95022-06010" in all_texts
    assert "90430-06817" in all_texts
    # Row 3 has NO valid qty in either model, so it MUST be omitted
    assert "99999-00000" not in all_texts
    assert "NEITHER MODEL" not in all_texts


def test_metadata_generate_and_export_endpoints():
    rows = [
        {
            "page": 7,
            "fig_no": "1",
            "fig_name": "CYLINDER HEAD",
            "ref_no": "1",
            "part_no": "BGP-E1111-00",
            "description": "HEAD, CYLINDER 1",
            "remarks": "",
            "BGPK": "1",
        }
    ]

    # 1. Test /api/meta/generate
    gen_res = client.post(
        "/api/meta/generate",
        json={
            "rows": rows,
            "model_columns": ["BGPK"],
            "brand": "YAMAHA",
            "model": "",
            "series": "series",
            "main_parts_only": True,
        },
    )
    assert gen_res.status_code == 200
    gen_data = gen_res.json()
    assert gen_data["total"] == 1
    item = gen_data["items"][0]
    assert item["fig_no"] == "1"
    assert item["part_name"] == "CYLINDER HEAD"
    assert item["image_filename"] == "YAM_BGPK_CYLINDER HEAD.jpeg"
    assert item["product_title"] == "YAMAHA, BGPK, CYLINDER HEAD"
    assert item["meta_short_description"] == "YAMAHA, BGPK, CYLINDER HEAD, INDIA SPARE"
    assert 120 <= len(item["meta_long_description"]) <= 140

    # Test /api/meta/generate with custom model typed after model code
    gen_custom = client.post(
        "/api/meta/generate",
        json={
            "rows": rows,
            "model_columns": ["BGPK"],
            "brand": "YAMAHA",
            "model": "R15",
            "series": "series",
            "main_parts_only": True,
        },
    )
    assert gen_custom.status_code == 200
    custom_item = gen_custom.json()["items"][0]
    assert custom_item["product_title"] == "YAMAHA, BGPK, R15, CYLINDER HEAD"
    assert custom_item["meta_short_description"] == "YAMAHA, BGPK, R15, CYLINDER HEAD, INDIA SPARE"
    assert 120 <= len(custom_item["meta_long_description"]) <= 140
    assert "BGPK R15" in custom_item["meta_long_description"]

    # 2. Test /api/meta/export as xlsx
    export_xlsx = client.post(
        "/api/meta/export",
        json={
            "items": gen_data["items"],
            "format": "xlsx",
            "filename": "Test_Meta",
        },
    )
    assert export_xlsx.status_code == 200
    assert "Test_Meta.xlsx" in export_xlsx.headers.get("content-disposition", "")
    assert export_xlsx.content[:4] == b"PK\x03\x04"

    # 3. Test /api/meta/export as csv
    export_csv = client.post(
        "/api/meta/export",
        json={
            "items": gen_data["items"],
            "format": "csv",
            "filename": "Test_Meta",
        },
    )
    assert export_csv.status_code == 200
    assert "Test_Meta.csv" in export_csv.headers.get("content-disposition", "")
    csv_txt = export_csv.content.decode("utf-8-sig")
    assert "Fig No.,Main Part Name,Brand,Model Code" in csv_txt
    assert "1,CYLINDER HEAD,YAMAHA,BGPK,,series,YAM_BGPK_CYLINDER HEAD.jpeg" in csv_txt



